import hashlib
import logging
import re
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


# ---------------------------------------------------------
# 헤더 후보
# ---------------------------------------------------------

HEADER_ALIASES = {
    "구분": ["구분", "구 분", "분류", "분류명", "category", "group"],
    "모델": ["모델명", "모델", "model", "model name"],
    "코드": [
        "code",
        "part no",
        "part number",
        "partno",
        "품번",
        "자산번호",
        "부품번호",
        "코드",
    ],
    "용량": ["용량", "capacity", "size"],
    "rpm": ["rpm", "회전속도"],
    "type": ["type", "타입", "종류", "방식"],
    "수량": [
        "수량",
        "재고수량",
        "재 고 수량",
        "qty",
        "quantity",
        "총 수량",
        "총수량",
        "재고",
    ],
    "위치": ["위치", "보관위치", "location", "storage"],
    "비고": ["비고", "비 고", "remark", "remarks", "note", "notes"],
}


# ---------------------------------------------------------
# 기본 정리
# ---------------------------------------------------------


def clean(value):
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    text = str(value).strip()

    if not text:
        return None

    return text


def normalize_header(value):
    if value is None:
        return ""

    return re.sub(r"[^0-9a-z가-힣]", "", str(value).lower())


# ---------------------------------------------------------
# Excel Matrix
# ---------------------------------------------------------


def read_matrix(ws):
    """
    시트 전체를 2차원 리스트로. ws.cell() 을 셀마다 호출하는 대신 iter_rows 로 한 번에 읽는다
    (수만 셀짜리 시트에서 수십 배 빠르다). 행 길이는 max_column 으로 맞춘다.
    """
    width = ws.max_column or 0
    matrix = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        row = list(row)
        if len(row) < width:
            row.extend([None] * (width - len(row)))
        matrix.append(row)
    return matrix


def fill_merged_cells(ws, matrix):
    """
    병합 셀의 좌상단 값을 병합 영역 전체에 복사한다.

    기존 코드처럼 모든 빈 셀을 무조건 아래로 복사하지 않는다.
    그렇게 하면 실제 빈 데이터까지 이전 행 값으로 오염될 수 있다.
    """

    for merged in ws.merged_cells.ranges:
        anchor = matrix[merged.min_row - 1][merged.min_col - 1]

        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                matrix[r - 1][c - 1] = anchor

    return matrix


# ---------------------------------------------------------
# 헤더 판단
# ---------------------------------------------------------


def known_header_score(value):
    """
    알려진 헤더면 높은 점수를 준다.
    """

    if value is None:
        return 0

    text = normalize_header(value)

    if not text:
        return 0

    score = 0

    for aliases in HEADER_ALIASES.values():
        for alias in aliases:
            alias_norm = normalize_header(alias)

            if text == alias_norm:
                return 4

            if alias_norm and alias_norm in text:
                score = max(score, 2)

    return score


def row_nonempty_count(row):
    return sum(clean(value) is not None for value in row)


def data_density(matrix, header_row):
    """
    헤더 다음 데이터 영역에서
    같은 컬럼들이 실제로 사용되고 있는지 확인한다.
    """

    if header_row + 1 >= len(matrix):
        return 0

    end = min(len(matrix), header_row + 11)

    data_rows = matrix[header_row + 1 : end]

    if not data_rows:
        return 0

    nonempty = 0
    total = 0

    for row in data_rows:
        for value in row:
            total += 1

            if clean(value) is not None:
                nonempty += 1

    if total == 0:
        return 0

    return nonempty / total


def header_score(matrix, row_index):
    """
    특정 행이 헤더일 가능성을 계산한다.

    기존처럼 특정 컬럼 하나에 의존하지 않고

    1. 알려진 헤더 여부
    2. 헤더 컬럼 수
    3. 문자열 형태
    4. 아래쪽 데이터 존재 여부

    를 함께 본다.
    """

    row = matrix[row_index]

    values = [clean(v) for v in row]

    values = [v for v in values if v is not None]

    if len(values) < 2:
        return 0

    score = 0

    # 알려진 헤더
    known = sum(known_header_score(v) for v in values)

    score += known

    # 헤더가 여러 컬럼이면 가산
    score += min(len(values), 12) * 0.5

    # 바로 아래 데이터가 존재하면 가산
    density = data_density(matrix, row_index)

    score += density * 5

    # 일반적으로 헤더는 문자열 위주
    text_count = sum(isinstance(v, str) for v in values)

    if text_count >= 2:
        score += 2

    return score


def detect_header(matrix):
    """
    시트 전체에서 헤더 후보를 찾는다.

    기존의 '첫 20행' 제한을 없애고
    앞쪽 50행까지 검사한다.
    """

    max_scan = min(len(matrix), 50)

    candidates = []

    for row_index in range(max_scan):
        score = header_score(matrix, row_index)

        if score > 0:
            candidates.append((score, row_index))

    if not candidates:
        return None

    candidates.sort(key=lambda x: x[0], reverse=True)

    best_score, best_row = candidates[0]

    # 너무 빈약한 행은 헤더로 취급하지 않는다.
    if best_score < 5:
        return None

    return best_row


# ---------------------------------------------------------
# 헤더 이름 생성
# ---------------------------------------------------------


def make_headers(matrix, header_row):
    headers = []

    used = {}

    for column_index, value in enumerate(matrix[header_row], start=1):
        header = clean(value)

        if not header:
            header = f"열_{get_column_letter(column_index)}"

        base = header

        count = used.get(base, 0) + 1

        used[base] = count

        if count > 1:
            header = f"{base}_{count}"

        headers.append(header)

    return headers


# ---------------------------------------------------------
# 실제 사용 컬럼 찾기
# ---------------------------------------------------------


def detect_active_columns(matrix, header_row, headers):
    active = []

    data_end = min(len(matrix), header_row + 101)

    for column_index, header in enumerate(headers):
        values = [
            clean(matrix[row_index][column_index])
            for row_index in range(header_row + 1, data_end)
        ]

        has_data = any(value is not None for value in values)

        # 헤더가 존재하거나 아래에 데이터가 있으면 유지
        if not header.startswith("열_") or has_data:
            active.append(column_index)

    return active


# ---------------------------------------------------------
# 데이터 종료 행
# ---------------------------------------------------------


def detect_data_end(matrix, header_row, active_columns):
    """
    데이터가 끝난 지점을 자동으로 찾는다.

    기존처럼 무조건 worksheet 마지막 행까지
    읽지 않는다.
    """

    last_data_row = header_row

    empty_streak = 0

    for row_index in range(header_row + 1, len(matrix)):
        has_data = any(clean(matrix[row_index][c]) is not None for c in active_columns)

        if has_data:
            last_data_row = row_index
            empty_streak = 0

        else:
            empty_streak += 1

            # 연속 2행 이상 비어 있으면
            # 데이터가 끝났다고 판단
            if empty_streak >= 2:
                break

    return last_data_row


# ---------------------------------------------------------
# 행 데이터
# ---------------------------------------------------------


def extract_rows(matrix, header_row, headers, active_columns):
    end_row = detect_data_end(matrix, header_row, active_columns)

    rows = []

    for row_index in range(header_row + 1, end_row + 1):
        data = {}

        has_data = False

        for column_index in active_columns:
            header = headers[column_index]

            value = clean(matrix[row_index][column_index])

            data[header] = value

            if value is not None:
                has_data = True

        if not has_data:
            continue

        rows.append(
            {
                "row_number": row_index + 1,
                "data": data,
            }
        )

    return rows


# ---------------------------------------------------------
# Table
# ---------------------------------------------------------


def detect_table(matrix, header_row):
    headers = make_headers(matrix, header_row)

    active_columns = detect_active_columns(matrix, header_row, headers)

    if not active_columns:
        return None

    rows = extract_rows(matrix, header_row, headers, active_columns)

    if not rows:
        return None

    return {
        "headers": [headers[c] for c in active_columns],
        "cols": active_columns,
        "start": header_row + 1,
        "end": rows[-1]["row_number"] - 1,
        "rows": rows,
    }


# ---------------------------------------------------------
# Sheet 타입
# ---------------------------------------------------------


def is_probably_image_sheet(ws, matrix, header_row):
    """
    표 헤더가 없는 시트만 이미지/기타로 분류한다.

    단순히 merged cell 개수나 column 개수로
    image라고 판단하지 않는다.
    """

    if header_row is not None:
        return False

    nonempty = 0

    for row in matrix:
        for value in row:
            if clean(value) is not None:
                nonempty += 1

    # 거의 비어 있는 시트
    if nonempty <= 5:
        return True

    return False


# ---------------------------------------------------------
# Workbook Parser
# ---------------------------------------------------------


def parse_workbook(path):
    """
    Excel 구조를 자동 분석한다.

    - 시트 개수 변경
    - 컬럼 개수 변경
    - 컬럼 추가/삭제
    - 장애파트/총계 컬럼 제거
    - 헤더 위치 변경
    - 빈 행 추가
    - 병합 셀 변경

    등을 최대한 구조적으로 처리한다.

    원본 Excel 컬럼은 data_json에 그대로 보존한다.
    """

    wb = load_workbook(path, data_only=True)

    result = []

    for ws in wb.worksheets:
        matrix = read_matrix(ws)

        matrix = fill_merged_cells(ws, matrix)

        header_row = detect_header(matrix)

        if header_row is None:
            result.append(
                {
                    "sheet_name": ws.title,
                    "sheet_type": (
                        "image"
                        if is_probably_image_sheet(ws, matrix, header_row)
                        else "other"
                    ),
                    "header_row": None,
                    "tables": [],
                    "dimensions": (ws.max_row, ws.max_column),
                }
            )

            continue

        table = detect_table(matrix, header_row)

        if table is None:
            result.append(
                {
                    "sheet_name": ws.title,
                    "sheet_type": "other",
                    "header_row": header_row + 1,
                    "tables": [],
                    "dimensions": (ws.max_row, ws.max_column),
                }
            )

            continue

        result.append(
            {
                "sheet_name": ws.title,
                "sheet_type": "table",
                "header_row": header_row + 1,
                "tables": [
                    {
                        "headers": table["headers"],
                        # 실제 Excel 열 인덱스(0-based). 빈 열을 건너뛴 경우 headers 와
                        # 순번이 어긋나므로 export 시 반드시 이 값을 써야 한다 (B-15).
                        "cols": table["cols"],
                        "rows": table["rows"],
                    }
                ],
                "dimensions": (ws.max_row, ws.max_column),
            }
        )

    logger.info("Workbook parse completed | path=%s", path)
    return result


# ---------------------------------------------------------
# SHA256
# ---------------------------------------------------------


def file_sha256(path):
    h = hashlib.sha256()

    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)

    return h.hexdigest()

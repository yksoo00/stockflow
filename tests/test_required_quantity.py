"""
품목별 필수수량: 출고 후 재고가 필수수량보다 적어지면 필수수량까지 모자란 만큼 입고요청.

  필수 2: 1개 남으면 1개 요청, 0개 남으면 2개 요청
  필수 1: 1개 남으면 요청 없음, 0개 남으면 1개 요청
"""

import io

import pytest
import sqlalchemy as sa
from flask_migrate import upgrade
from openpyxl import load_workbook

from app import create_app, db
from app.models import ExcelFile, InventoryRow, SheetColumn, StockRequest
from app.services.inventory.normalize import normalize_row

from .conftest import api, make_row
from .test_excel import _upload, _xlsx


def _row_with_required(app, quantity, required, identifier="ST4000"):
    row_id = make_row(app, identifier=identifier, quantity=quantity)
    with app.app_context():
        row = db.session.get(InventoryRow, row_id)
        row.data_json = {**row.data_json, "필수수량": required}
        row.required_quantity = required
        db.session.commit()
    return row_id


def _auto_requests(app, row_id):
    with app.app_context():
        return [
            r.quantity
            for r in StockRequest.query.filter_by(inventory_row_id=row_id, source="auto")
            .order_by(StockRequest.id)
            .all()
        ]


def _out(client, row_id, quantity):
    resp = api(client, "POST", "/api/stockout", row_id=row_id, site="A", quantity=quantity)
    assert resp.status_code == 200, resp.get_json()
    return resp.get_json()


# ---------------------------------------------------------------------------
# 자동 입고요청 규칙
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "required, stock, out, expected",
    [
        (2, 2, 1, [1]),  # 필수 2, 1개 남음 → 1개
        (2, 2, 2, [2]),  # 필수 2, 0개 남음 → 2개
        (1, 2, 1, []),  # 필수 1, 1개 남음 → 없음
        (1, 1, 1, [1]),  # 필수 1, 0개 남음 → 1개
        (3, 5, 4, [2]),  # 필수 3, 1개 남음 → 2개
        (0, 1, 1, []),  # 필수 0 → 요청 안 함
    ],
)
def test_auto_request_fills_up_to_required(user_client, app, required, stock, out, expected):
    row_id = _row_with_required(app, stock, required)
    data = _out(user_client, row_id, out)
    assert data["auto_request_created"] is bool(expected)
    assert _auto_requests(app, row_id) == expected


def test_empty_required_falls_back_to_two(user_client, app):
    """필수수량 칸이 비어 있으면 예전처럼 2개 기준."""
    row_id = make_row(app, quantity=2)
    _out(user_client, row_id, 1)
    assert _auto_requests(app, row_id) == [1]


def test_in_flight_requests_are_not_requested_twice(user_client, app):
    """도착 전인 요청 수량은 빼고 요청한다 — 합계가 필수수량을 넘지 않는다."""
    row_id = _row_with_required(app, 2, 2)
    _out(user_client, row_id, 1)  # 1개 남음 → 1개 요청
    _out(user_client, row_id, 1)  # 0개 남음 → 필요 2, 진행 중 1 → 1개만 추가
    assert _auto_requests(app, row_id) == [1, 1]


def test_arrived_request_no_longer_counts_as_in_flight(admin_client, app):
    row_id = _row_with_required(app, 2, 2)
    _out(admin_client, row_id, 1)
    with app.app_context():
        rid = StockRequest.query.filter_by(inventory_row_id=row_id).one().id
    api(admin_client, "POST", f"/api/stockrequest/{rid}/approve")
    api(admin_client, "POST", f"/api/stockrequest/{rid}/arrive")  # 재고 2개로 복귀

    _out(admin_client, row_id, 2)  # 0개 남음, 진행 중 없음 → 2개
    assert _auto_requests(app, row_id) == [1, 2]


def test_editing_required_in_sheet_changes_rule(admin_client, app):
    """시트에서 필수수량을 고치면 다음 출고부터 그 값으로 채운다."""
    row_id = _row_with_required(app, 3, 2)
    resp = api(admin_client, "PATCH", f"/api/rows/{row_id}", data={"필수수량": "3"})
    assert resp.status_code == 200, resp.get_json()
    with app.app_context():
        assert db.session.get(InventoryRow, row_id).required_quantity == 3

    _out(admin_client, row_id, 1)  # 2개 남음 < 필수 3 → 1개
    assert _auto_requests(app, row_id) == [1]


def test_low_stock_list_uses_required(user_client, app):
    below = _row_with_required(app, 2, 3, identifier="BELOW")  # 2 < 3 → 저재고
    _row_with_required(app, 1, 1, identifier="EXACT")  # 1 == 1 → 정상
    _row_with_required(app, 5, 2, identifier="PLENTY")
    data = user_client.get("/api/search?low=1").get_json()
    assert [item["id"] for item in data["items"]] == [below]


# ---------------------------------------------------------------------------
# 엑셀 업로드 / 화면 / 내보내기
# ---------------------------------------------------------------------------
def test_required_header_is_not_mistaken_for_stock():
    n = normalize_row({"재 고": 4, "필수수량": "2"})
    assert n["quantity"] == 4 and n["required_quantity"] == 2
    assert "required_quantity" not in normalize_row({"수량": 1, "필수수량": ""})


def test_upload_with_required_column_uses_its_values(admin_client, app):
    rows = [
        ["Code", "품명", "수량", "용량", "위치", "필수수량"],
        ["A1", "디스크", 5, "4TB", "1-1", 1],
        ["A2", "디스크", 5, "4TB", "1-1", 3],
        ["A3", "디스크", 5, "4TB", "1-1", None],  # 빈 칸 → 2
    ]
    _upload(admin_client, _xlsx(rows))
    with app.app_context():
        # 파서는 셀 값을 문자열로 보관한다 (수량도 마찬가지)
        got = {r.identifier: (r.required_quantity, str(r.data_json["필수수량"])) for r in InventoryRow.query}
        assert got == {"A1": (1, "1"), "A2": (3, "3"), "A3": (2, "2")}
        assert SheetColumn.query.filter_by(original_name="필수수량").count() == 1  # 새로 만들지 않음

    columns = admin_client.get("/api/search").get_json()["columns"]
    assert columns[-1] == "필수수량"


def test_upload_without_required_column_adds_it_and_exports_it(admin_client, app):
    rows = [["Code", "품명", "수량", "용량", "위치"], ["A1", "디스크", 5, "4TB", "1-1"]]
    _upload(admin_client, _xlsx(rows, title_rows=[["재고 현황"]]))
    with app.app_context():
        row = InventoryRow.query.one()
        assert row.required_quantity == 2 and row.data_json["필수수량"] == 2
        file_id, row_id = row.excel_file_id, row.id

    # 화면에서 1 로 고친 뒤 내보내면 원본 맨 끝 열(F)에 헤더와 값이 들어간다
    api(admin_client, "PATCH", f"/api/rows/{row_id}", data={"필수수량": 1})
    ws = load_workbook(io.BytesIO(admin_client.get(f"/files/{file_id}/export").data))["재고"]
    assert ws["A1"].value == "재고 현황"  # 원본 서식 경로
    assert ws["F2"].value == "필수수량"
    assert ws["F3"].value == 1


# ---------------------------------------------------------------------------
# 마이그레이션: 이미 올라와 있는 시트에 필수수량 열 채우기
# ---------------------------------------------------------------------------
def test_migration_backfills_existing_sheets(tmp_path, monkeypatch):
    monkeypatch.setenv("AUTO_MIGRATE", "0")
    monkeypatch.setenv("LOG_DIR", str(tmp_path))
    app = create_app(
        test_config={
            "TESTING": False,
            "SECRET_KEY": "x" * 32,
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'm.db'}",
        }
    )
    with app.app_context():
        upgrade(revision="0003_user_sites_text")
        conn = db.session.connection()
        conn.execute(sa.text(
            "INSERT INTO excel_file (id, original_filename, stored_filename, file_path) "
            "VALUES (1, 'a.xlsx', 'a.xlsx', '/x/a.xlsx')"
        ))
        conn.execute(sa.text(
            "INSERT INTO excel_sheet (id, excel_file_id, sheet_name) VALUES (1, 1, 'S1'), (2, 1, 'S2')"
        ))
        conn.execute(sa.text(
            "INSERT INTO detected_table (id, sheet_id, header_row) VALUES (1, 1, 1), (2, 2, 1)"
        ))
        # S1: 필수수량 열 없음 / S2: 엑셀에 이미 있음 (한 칸은 비어 있음)
        conn.execute(sa.text(
            "INSERT INTO sheet_column (sheet_id, table_id, column_index, original_name) VALUES "
            "(1, 1, 1, 'Code'), (1, 1, 2, '수량'), "
            "(2, 2, 1, 'Code'), (2, 2, 2, '수량'), (2, 2, 3, '필수수량')"
        ))
        conn.execute(sa.text(
            "INSERT INTO inventory_row (id, excel_file_id, sheet_id, table_id, row_number, data_json, is_deleted) VALUES "
            "(1, 1, 1, 1, 2, '{\"Code\": \"A\", \"수량\": 3}', 0), "
            "(2, 1, 2, 2, 2, '{\"Code\": \"B\", \"수량\": 3, \"필수수량\": 1}', 0), "
            "(3, 1, 2, 2, 3, '{\"Code\": \"C\", \"수량\": 3, \"필수수량\": \"\"}', 0)"
        ))
        db.session.commit()

        upgrade()

        s1_cols = {c.original_name: c.column_index for c in SheetColumn.query.filter_by(sheet_id=1)}
        assert s1_cols == {"Code": 1, "수량": 2, "필수수량": 3}
        assert SheetColumn.query.filter_by(sheet_id=2, original_name="필수수량").count() == 1
        rows = {r.id: (r.required_quantity, r.data_json.get("필수수량")) for r in InventoryRow.query}
        assert rows == {1: (2, 2), 2: (1, 1), 3: (2, 2)}
        assert db.session.get(ExcelFile, 1) is not None
